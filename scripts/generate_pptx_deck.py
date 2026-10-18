#!/usr/bin/env python3
"""PayLens PowerPoint Pitch Deck Generator.

Generates 16:9 widescreen presentation docs/paylens_pitch_deck.pptx
from verified 1080p slide renders in docs/slides/.
"""

import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

ROOT_DIR = Path(__file__).resolve().parent.parent
DOCS_DIR = ROOT_DIR / "docs"
SLIDES_DIR = DOCS_DIR / "slides"
OUTPUT_PPTX = DOCS_DIR / "paylens_pitch_deck.pptx"

def main():
    print("=" * 70)
    print("PAYLENS POWERPOINT PITCH DECK GENERATOR")
    print("=" * 70)

    # 1. Ensure slide PNGs exist
    slide_files = [SLIDES_DIR / f"slide_{i:02d}.png" for i in range(1, 8)]
    missing = [f for f in slide_files if not f.exists()]
    
    if missing:
        print(f"Missing {len(missing)} slide image(s). Generating slides first...")
        import subprocess
        subprocess.run([sys.executable, str(ROOT_DIR / "scripts" / "generate_pitch_deck.py")], check=True)

    # 2. Initialize 16:9 presentation
    prs = Presentation()
    # 16:9 aspect ratio standard dimensions: 13.333 x 7.5 inches
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]  # completely blank slide

    # 3. Add each 1920x1080 slide image full-bleed
    for i, slide_path in enumerate(slide_files, start=1):
        if not slide_path.exists():
            print(f"[ERROR] Expected slide file not found: {slide_path}")
            sys.exit(1)
            
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(slide_path),
            Inches(0),
            Inches(0),
            width=prs.slide_width,
            height=prs.slide_height
        )
        print(f"  - Embedded Slide {i:02d}/07: {slide_path.name}")

    # 4. Save presentation
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PPTX))

    file_size_kb = os.path.getsize(OUTPUT_PPTX) / 1024
    print(f"\n[SUCCESS] Saved presentation to {OUTPUT_PPTX} ({file_size_kb:.1f} KB, 7 slides)")

if __name__ == "__main__":
    main()
